"""Usage: analyze_excel.py <file.xlsm>
Prints structural metadata (sheet names, sizes, headers, column types) and a Gemma-written anonymized description.
Raw cell values are only sent to the local Ollama model, never printed."""
import re
import sys
import zipfile
from collections import Counter

import openpyxl

from llm import ask

path = sys.argv[1]
wb_f = openpyxl.load_workbook(path, data_only=False)
wb_v = openpyxl.load_workbook(path, data_only=True)

with zipfile.ZipFile(path) as z:
    vba = [n for n in z.namelist() if "vbaProject" in n]
print(f"# Classeur : {len(wb_f.sheetnames)} feuilles, macros VBA : {'oui' if vba else 'non'}\n")


def ctype(v):
    if v is None:
        return None
    return type(v).__name__


def norm_formula(f):
    return re.sub(r"(\$?[A-Z]{1,3}\$?)\d+", r"\1#", f)


for name in wb_f.sheetnames:
    ws_f, ws_v = wb_f[name], wb_v[name]
    rows_v = list(ws_v.iter_rows(values_only=True))
    rows_f = list(ws_f.iter_rows(values_only=True))
    nonempty = [r for r in rows_v if any(c not in (None, "") for c in r)]
    formulas = Counter(norm_formula(c) for r in rows_f for c in r if isinstance(c, str) and c.startswith("="))

    header_idx = next((i for i, r in enumerate(rows_v[:15]) if sum(isinstance(c, str) and c.strip() != "" for c in r) >= 3), None)
    print(f"## Feuille « {name} » — {ws_v.max_row} lignes x {ws_v.max_column} col., {len(nonempty)} lignes non vides, {sum(formulas.values())} formules")
    if header_idx is not None:
        header = rows_v[header_idx]
        print(f"En-têtes (ligne {header_idx + 1}) :")
        for j, h in enumerate(header):
            if h in (None, ""):
                continue
            col = [r[j] for r in rows_v[header_idx + 1:] if j < len(r)]
            types = Counter(t for t in map(ctype, col) if t)
            print(f"  - {h!s} : {sum(types.values())} valeurs, types {dict(types)}")
    if formulas:
        print("Formules types :", "; ".join(f"{f} (x{n})" for f, n in formulas.most_common(6)))

    if not nonempty:
        print()
        continue
    start = header_idx or 0
    sample = rows_v[start:start + 45]
    tsv = "\n".join("\t".join("" if c is None else str(c) for c in r) for r in sample)
    fsample = "\n".join(f"{f}" for f, _ in formulas.most_common(15))
    print("\n" + ask(
        f"Voici un extrait (TSV) de la feuille « {name} » d'un classeur de comptabilité d'une ostéopathe libérale "
        f"(déclaration contrôlée BNC, formulaire 2035).\n\nFormules les plus fréquentes :\n{fsample}\n\nExtrait :\n{tsv}\n\n"
        "Décris : 1) le rôle de la feuille ; 2) chaque colonne (signification, type, format, valeurs possibles si liste fermée non personnelle) ; "
        "3) les calculs/totaux ; 4) ce qu'une application de gestion devrait reprendre ou importer de cette feuille."
    ) + "\n")

if vba:
    try:
        from oletools.olevba import VBA_Parser
        p = VBA_Parser(path)
        code = "\n\n".join(f"' --- {vn}\n{src}" for _, _, vn, src in p.extract_macros() if src.strip())[:40000]
        if code:
            print("## Macros VBA\n" + ask(
                "Voici le code VBA d'un classeur de comptabilité d'ostéopathe. Décris fonctionnellement ce que fait chaque macro "
                "(ex. génère une facture, numérote, exporte) sans recopier de données personnelles :\n\n" + code))
    except ImportError:
        print("(oletools absent : macros non analysées)")
